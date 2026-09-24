#!/usr/bin/env python3
"""Import checkpoint generation sweeps and long-run validation loss into Excel."""

from __future__ import annotations

import argparse
import json
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "results" / "Results_accuracy-125-partial_20260921.xlsx"
DEFAULT_OUTPUT = ROOT / "results" / "Results_training-progression-long32i-64i_20260923.xlsx"
DRIVE_ROOT = (
    Path.home()
    / "Library/CloudStorage/GoogleDrive-ruurd.kuiper@gmail.com/My Drive/lad-generic-results"
)
RUN_IDS = {
    32: "20260923T134029.046940Z--colab-validation",
    64: "20260923T190517.263478Z--colab-validation",
    128: "20260923T201036.622289Z--colab-validation",
}
ALL_GLOBAL_STEPS = (1_000, 5_000, 10_000, 15_000, 20_000, 25_000, 30_000, 35_000, 40_000, 45_000, 50_000)
SELECTED_GLOBAL_STEPS = {
    32: ALL_GLOBAL_STEPS,
    64: ALL_GLOBAL_STEPS,
    128: ALL_GLOBAL_STEPS,
}
METRIC_FIELDS = {
    "perplexity": "perplexity",
    "distinct_1": "mean_distinct_1",
    "distinct_2": "mean_distinct_2",
    "distinct_3": "mean_distinct_3",
}


def _global_step(model: str) -> int | None:
    try:
        local_step = int(model.rsplit("checkpoint-", 1)[1])
    except (IndexError, ValueError):
        return None
    if model.startswith("llama-3.1-8b-mask-continued/"):
        return 25_000 + local_step
    if model.startswith(("llama-3.1-8b-mask/", "llama-3.1-8b-mask-long/")):
        return local_step
    return None


def _load_sweep(path: Path, expected_nfe: int) -> tuple[dict[int, dict[str, float]], dict]:
    run = json.loads((path / "run.json").read_text())
    summary = json.loads((path / "summary.json").read_text())
    if run.get("status") != "completed":
        raise ValueError(f"Benchmark run is not complete: {path}")
    settings = run["config"]["mask_only_task_generation"]["open_ended"]
    if int(settings["num_steps"]) != expected_nfe or int(settings["max_new_tokens"]) != 128:
        raise ValueError(f"Unexpected generation settings in {path}: {settings}")

    rows: dict[int, dict[str, float]] = {}
    for model_entry in summary["models"]:
        for result in model_entry["results"]:
            if result.get("task") != "open_ended" or result.get("method") != "diffusion":
                continue
            step = _global_step(str(result["model"]))
            if step not in SELECTED_GLOBAL_STEPS[expected_nfe]:
                continue
            if int(result["inference_settings"]["num_steps"]) != expected_nfe:
                raise ValueError(f"Result NFE mismatch for {result['model']}")
            rows[step] = {
                name: float(result[field]) for name, field in METRIC_FIELDS.items()
            }
            rows[step]["total"] = int(result["total"])
    missing = set(SELECTED_GLOBAL_STEPS[expected_nfe]) - set(rows)
    if missing:
        raise ValueError(f"Missing {expected_nfe}-NFE checkpoints: {sorted(missing)}")
    return rows, {"run": run, "summary": summary, "path": path}


def _copy_cell_style(source, destination) -> None:
    destination._style = copy(source._style)
    destination.number_format = source.number_format
    destination.alignment = copy(source.alignment)
    destination.protection = copy(source.protection)


def _copy_block_style(sheet, source_row: int, target_row: int) -> None:
    for offset in range(5):
        sheet.row_dimensions[target_row + offset].height = sheet.row_dimensions[source_row + offset].height
        for column in range(1, sheet.max_column + 1):
            _copy_cell_style(sheet.cell(source_row + offset, column), sheet.cell(target_row + offset, column))


def _write_progression_block(
    sheet,
    *,
    header_row: int,
    template_row: int,
    metric: str,
    title: str,
    sweeps: dict[int, dict[int, dict[str, float]]],
    baselines: dict[int, float],
    autoregressive: float,
    preserve_128: bool = False,
) -> None:
    if header_row != template_row:
        _copy_block_style(sheet, template_row, header_row)
    step_columns = {4 + index: step for index, step in enumerate((1_000, 5_000, 10_000, 15_000, 20_000, 25_000, 30_000, 35_000, 40_000, 45_000, 50_000))}
    sheet.cell(header_row, 3, "Training iterations")
    for column, step in step_columns.items():
        sheet.cell(header_row, column, step)
    sheet.cell(header_row + 1, 3, title)

    for row_offset, nfe in enumerate((32, 64, 128), start=2):
        row = header_row + row_offset
        sheet.cell(row, 3, f"BYOD-Llama ({nfe} it.)")
        if not (preserve_128 and nfe == 128):
            for column in step_columns:
                sheet.cell(row, column, None)
        if nfe in sweeps:
            for column, step in step_columns.items():
                if step in sweeps[nfe]:
                    sheet.cell(row, column, sweeps[nfe][step][metric])
        sheet.cell(row, 15, f"LLaDA ({nfe} it.)")
        sheet.cell(row, 16, baselines[nfe])
        sheet.cell(row, 17, "Llama 3.1 8B AR")
        sheet.cell(row, 18, autoregressive)


def _autoregressive_reference(metadata: dict, metric: str) -> tuple[float, int, str]:
    """Extract the single Llama AR reference from the 32-NFE benchmark run."""
    matches = [
        result
        for model in metadata["summary"]["models"]
        for result in model["results"]
        if result.get("task") == "open_ended"
        and result.get("method") == "autoregressive"
        and result.get("model") == "meta-llama/Llama-3.1-8B-Instruct"
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one Llama AR open-ended result, found {len(matches)}")
    result = matches[0]
    return float(result[METRIC_FIELDS[metric]]), int(result["total"]), str(result["model"])


def _append_provenance(sheet, rows: list[tuple[str, object]]) -> None:
    row = sheet.max_row + 2
    for field, value in rows:
        sheet.cell(row, 1, field)
        sheet.cell(row, 2, value)
        row += 1


def _write_validation_sheet(workbook: Workbook, metrics_path: Path) -> tuple[int, int]:
    title = "Long validation loss"
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    records = [
        json.loads(line)
        for line in metrics_path.read_text().splitlines()
        if line.strip()
    ]
    records = [record for record in records if record.get("split") == "validation"]
    records.sort(key=lambda record: int(record["step"]))
    if not records:
        raise ValueError(f"No validation records in {metrics_path}")

    headers = (
        "Training update",
        "Weighted validation loss",
        "Unweighted masked-token CE",
        "Generation perplexity (when measured)",
        "Mean Distinct-1 (when measured)",
    )
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row, record in enumerate(records, start=2):
        values = (
            int(record["step"]),
            float(record["weighted_loss"]),
            float(record["unweighted_masked_token_ce"]),
            record.get("generation_perplexity"),
            record.get("generation_mean_distinct_1"),
        )
        for column, value in enumerate(values, start=1):
            sheet.cell(row, column, value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    for column, width in zip("ABCDE", (18, 25, 28, 38, 34)):
        sheet.column_dimensions[column].width = width
    for row in range(2, sheet.max_row + 1):
        for column in range(2, 6):
            sheet.cell(row, column).number_format = "0.000000"
    return int(records[0]["step"]), int(records[-1]["step"])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--drive-root", type=Path, default=DRIVE_ROOT)
    args = parser.parse_args()

    sweeps: dict[int, dict[int, dict[str, float]]] = {}
    metadata = {}
    for nfe, run_id in RUN_IDS.items():
        run_path = args.drive_root / "validation/benchmark_runs" / run_id
        sweeps[nfe], metadata[nfe] = _load_sweep(run_path, nfe)

    workbook = load_workbook(args.input)
    primary = workbook["Primary outcomes"]
    training = workbook["Training progresion"]

    perplexity_baselines = {32: float(primary["P18"].value), 64: float(primary["P19"].value), 128: float(primary["P20"].value)}
    distinct_baselines = {
        "distinct_1": {32: float(primary["P23"].value), 64: float(primary["P24"].value), 128: float(primary["P25"].value)},
        "distinct_2": {32: float(primary["P28"].value), 64: float(primary["P29"].value), 128: float(primary["P30"].value)},
        "distinct_3": {32: float(primary["P33"].value), 64: float(primary["P34"].value), 128: float(primary["P35"].value)},
    }
    ar_references = {
        metric: _autoregressive_reference(metadata[32], metric)[0]
        for metric in METRIC_FIELDS
    }
    ar_value, ar_total, ar_model = _autoregressive_reference(metadata[32], "perplexity")

    _write_progression_block(primary, header_row=42, template_row=42, metric="perplexity", title="128 tokens generated — Phi-4 token-weighted perplexity", sweeps=sweeps, baselines=perplexity_baselines, autoregressive=ar_references["perplexity"], preserve_128=True)
    for header_row, metric, title in (
        (49, "distinct_1", "128 tokens generated — mean sliding model-token Distinct-1"),
        (56, "distinct_2", "128 tokens generated — mean sliding model-token Distinct-2"),
        (63, "distinct_3", "128 tokens generated — mean sliding model-token Distinct-3"),
    ):
        _write_progression_block(primary, header_row=header_row, template_row=42, metric=metric, title=title, sweeps=sweeps, baselines=distinct_baselines[metric], autoregressive=ar_references[metric])

    _write_progression_block(training, header_row=4, template_row=4, metric="perplexity", title="128 tokens generated — Phi-4 token-weighted perplexity", sweeps=sweeps, baselines=perplexity_baselines, autoregressive=ar_references["perplexity"], preserve_128=True)
    for header_row, metric, title in (
        (17, "distinct_1", "128 tokens generated — mean sliding model-token Distinct-1"),
        (24, "distinct_2", "128 tokens generated — mean sliding model-token Distinct-2"),
        (31, "distinct_3", "128 tokens generated — mean sliding model-token Distinct-3"),
    ):
        _write_progression_block(training, header_row=header_row, template_row=4, metric=metric, title=title, sweeps=sweeps, baselines=distinct_baselines[metric], autoregressive=ar_references[metric])

    metrics_path = args.drive_root / "outputs/llama-3.1-8b-mask-long/metrics.jsonl"
    first_validation_step, last_validation_step = _write_validation_sheet(workbook, metrics_path)

    provenance = workbook["Import provenance"]
    provenance_rows: list[tuple[str, object]] = []
    for nfe in (32, 64, 128):
        run = metadata[nfe]["run"]
        provenance_rows.extend((
            (f"{nfe}-NFE progression run ID", run["run_id"]),
            (f"{nfe}-NFE progression completed UTC", run["completed_at"]),
            (f"{nfe}-NFE progression source", str(metadata[nfe]["path"] / "summary.json")),
            (f"{nfe}-NFE imported global updates", ", ".join(f"{step // 1000}k" for step in SELECTED_GLOBAL_STEPS[nfe])),
            (f"{nfe}-NFE scored prompts per checkpoint", sweeps[nfe][1_000]["total"]),
            (f"{nfe}-NFE imported metrics", "Token-weighted perplexity; mean sliding model-token Distinct-1/2/3"),
        ))
    provenance_rows.extend((
        ("Autoregressive reference run ID", metadata[32]["run"]["run_id"]),
        ("Autoregressive reference model", ar_model),
        ("Autoregressive reference scored prompts", ar_total),
        ("Autoregressive reference perplexity", ar_value),
        ("Autoregressive reference mean Distinct-1", ar_references["distinct_1"]),
        ("Long-run validation-loss source", str(metrics_path)),
        ("Long-run validation-loss snapshot", f"steps {first_validation_step}--{last_validation_step} in increments of 500"),
        ("Progression import timestamp UTC", datetime.now(timezone.utc).isoformat()),
    ))
    _append_provenance(provenance, provenance_rows)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Wrote {args.output}")
    print(f"Long-run validation-loss snapshot ends at step {last_validation_step:,}")


if __name__ == "__main__":
    main()
